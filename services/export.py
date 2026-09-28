"""Экспорт данных в Excel (openpyxl, без pandas).

Три листа: «Яндекс 360», «AD», «Расхождения». Возвращает байты .xlsx —
UI отдаёт их через ui.download, ядро не зависит от фреймворка.
"""
from __future__ import annotations

from datetime import datetime
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from core.logs import get_logger

logger = get_logger(__name__)

_HEADER_FILL = PatternFill("solid", fgColor="1F2937")
_HEADER_FONT = Font(color="FFFFFF", bold=True)


def _write_sheet(wb: Workbook, title: str, headers: list[str], rows: list[list], first: bool = False) -> None:
    ws = wb.active if first else wb.create_sheet()
    ws.title = title
    ws.append(headers)
    for cell in ws[1]:
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(vertical="center")
    for row in rows:
        ws.append(row)
    # Автоширина по содержимому (с разумным ограничением)
    for col_idx, header in enumerate(headers, start=1):
        width = max(
            [len(str(header))] + [len(str(r[col_idx - 1])) for r in rows if len(r) >= col_idx] or [10]
        )
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max(width + 2, 10), 60)
    ws.freeze_panes = "A2"


def yandex_rows(yandex_users: list) -> list[list]:
    return [
        [u.nickname, u.display_name, u.department_name, u.position, u.email,
         u.external_id or "", ", ".join(u.phones),
         "Да" if u.is_enabled else "Нет", "Да" if u.is_admin else "Нет"]
        for u in yandex_users
    ]


def ad_rows(ad_users: list) -> list[list]:
    return [
        [u.sam_account_name, u.display_name, u.department, u.company, u.email_address,
         u.snils or "", u.office_phone, u.mobile_phone,
         "Да" if u.enabled else "Нет"]
        for u in ad_users
    ]


def mismatch_rows(comparison: dict | None) -> list[list]:
    """Плоские строки расхождений: одна строка на пару поле AD/Яндекс."""
    rows: list[list] = []
    for item in (comparison or {}).get("mismatched", []):
        for diff in item.get("differences", []):
            rows.append([
                item.get("login", ""),
                diff.get("label", diff.get("field", "")),
                diff.get("ad", ""),
                diff.get("yandex", ""),
            ])
    return rows


def build_report(yandex_users: list, ad_users: list, comparison: dict | None = None) -> bytes:
    """Собирает .xlsx с пользователями Яндекс/AD и расхождениями. Возвращает байты."""
    wb = Workbook()
    _write_sheet(
        wb, "Яндекс 360",
        ["Логин", "Имя", "Отдел", "Должность", "Email", "External ID (СНИЛС)",
         "Телефоны", "Активен", "Админ"],
        yandex_rows(yandex_users), first=True,
    )
    _write_sheet(
        wb, "AD",
        ["Логин", "Имя", "Отдел", "Компания", "Email", "СНИЛС",
         "Рабочий телефон", "Мобильный телефон", "Активен"],
        ad_rows(ad_users),
    )
    _write_sheet(
        wb, "Расхождения",
        ["Логин", "Поле", "AD", "Яндекс 360"],
        mismatch_rows(comparison),
    )
    buffer = BytesIO()
    wb.save(buffer)
    logger.info("Экспорт Excel собран: %s Яндекс / %s AD", len(yandex_users), len(ad_users))
    return buffer.getvalue()


def report_filename() -> str:
    return f"y360-report_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
